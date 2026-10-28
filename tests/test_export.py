import json

from openpyxl import load_workbook

from tests.conftest import VALID_PROPOSAL, WORKSPACE


def test_export_marks_only_confirmed_entries(run, analyzed, project):
    run("propose", WORKSPACE, "STAT", stdin=VALID_PROPOSAL)
    run("decide", WORKSPACE, "STAT", "--status", "übernommen")
    run("propose", WORKSPACE, "MAT_KAT", stdin={
        **VALID_PROPOSAL, "beschreibung": "Materialkategorie des Auftrags.", "lesbarer_name": "Materialkategorie",
        "konfidenz": "mittel", "offene_fragen": [{"frage": "Welche Schreibweise ist verbindlich?", "ansprechperson": "Einkauf"}],
    })
    code, result, _ = run("export", WORKSPACE)
    assert code == 0
    assert result["offene_fragen"] == 1
    assert any("noch nicht geprüft" in w for w in result["warnungen"])

    folder = project / "ausgabe" / WORKSPACE
    assert {p.name for p in folder.iterdir()} >= {
        "datenwoerterbuch.md", "datenwoerterbuch.xlsx", "datasheet.md", "schema.json", "offene_fragen.md",
    }
    dictionary = (folder / "datenwoerterbuch.md").read_text(encoding="utf-8")
    stat_row = next(line for line in dictionary.splitlines() if line.startswith("| STAT |"))
    mat_row = next(line for line in dictionary.splitlines() if line.startswith("| MAT_KAT |"))
    assert "geprüft (übernommen)" in stat_row
    assert "Vorschlag – nicht bestätigt" in mat_row and "geprüft" not in mat_row
    assert dictionary.count("geprüft (") == 1

    questions = (folder / "offene_fragen.md").read_text(encoding="utf-8")
    assert "Einkauf" in questions and "Welche Schreibweise" in questions

    schema = json.loads((folder / "schema.json").read_text(encoding="utf-8"))
    assert len(schema["fields"]) == 11
    stat = next(f for f in schema["fields"] if f["name"] == "STAT")
    assert stat["x-pruefstatus"] == "geprüft (übernommen)" and stat["title"] == "Auftragsstatus"
    assert next(f for f in schema["fields"] if f["name"] == "DAT_ERF")["format"] == "any"

    workbook = load_workbook(folder / "datenwoerterbuch.xlsx")
    assert workbook.sheetnames == ["Spalten", "Befunde", "Offene Fragen"]
    assert workbook["Spalten"].max_row == 12


def test_datasheet_export_has_all_sections(run, analyzed, project):
    run("datasheet", WORKSPACE, stdin={"zweck": "Auftragsverwaltung", "weitergabe": "nur intern"})
    run("datasheet", WORKSPACE, "--status", "übernommen")
    run("export", WORKSPACE)
    sheet = (project / "ausgabe" / WORKSPACE / "datasheet.md").read_text(encoding="utf-8")
    for heading in ("Motivation", "Zusammensetzung", "Erhebung", "Vorverarbeitung", "Einschränkungen", "Weitergabe", "Pflege"):
        assert heading in sheet
    assert "geprüft (übernommen)" in sheet and "nur intern" in sheet
    assert "KD_NAME" in sheet  # listed as personal column


def test_excel_export_never_contains_live_formulas(run, analyzed, project):
    # Text from the model or from data cells may start with "=": it must stay text in Excel.
    formula = '=HYPERLINK("http://example.invalid","Klick")'
    run("propose", WORKSPACE, "STAT", stdin={**VALID_PROPOSAL, "beschreibung": formula, "wertebereich": "=1+1"})
    assert run("export", WORKSPACE)[0] == 0

    workbook = load_workbook(project / "ausgabe" / WORKSPACE / "datenwoerterbuch.xlsx")
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    texts = [c.value for c in cells if isinstance(c.value, str)]
    assert formula in texts and "=1+1" in texts
