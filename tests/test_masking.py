"""No clear personal value may leave the core: not on stdout, not in any exported file."""

import re

import pytest
from openpyxl import load_workbook

from datadoc import pii
from tests.conftest import DEMO, VALID_PROPOSAL, WORKSPACE


def raw_names(rows):
    return {re.sub(r"^Dr\.\s+", "", r["KD_NAME"]) for r in rows}


def raw_phones(rows):
    phones = {r["KD_TEL"] for r in rows}
    for r in rows:
        phones |= {m.group(0) for m in pii.PHONE_RE.finditer(r["BEM"])}
    return phones


def subscriber_numbers(phones):
    """The last digit group of each phone number, e.g. ``123456`` in ``07131/123456``."""
    return {re.split(r"[ \-/]", p)[-1] for p in phones}


def everything_written(run, project) -> str:
    texts = list(run.outputs)
    for path in (project / "ausgabe").rglob("*"):
        if path.suffix in (".md", ".json"):
            texts.append(path.read_text(encoding="utf-8"))
        elif path.suffix == ".xlsx":
            workbook = load_workbook(path)
            texts += [str(c.value) for ws in workbook for row in ws.iter_rows() for c in row if c.value is not None]
    return "\n".join(texts)


@pytest.fixture
def full_session(run, demo_rows):
    run("analyze", DEMO)
    for column in demo_rows[0]:
        run("column", WORKSPACE, column)
    run("propose", WORKSPACE, "STAT", stdin=VALID_PROPOSAL)
    run("decide", WORKSPACE, "STAT", "--status", "übernommen")
    run("datasheet", WORKSPACE, stdin={"zweck": "Auftragsverwaltung", "ansprechperson": "Büro"})
    run("status", WORKSPACE)
    run("export", WORKSPACE)
    return run


def test_no_clear_names_anywhere(full_session, project, demo_rows):
    text = everything_written(full_session, project)
    leaked = sorted(n for n in raw_names(demo_rows) if n in text)
    assert leaked == []
    assert not re.search(r"\b(Herr|Frau)\s+(?!\[NAME\])[A-ZÄÖÜ]", text)


def test_no_phone_numbers_anywhere(full_session, project, demo_rows):
    text = everything_written(full_session, project)
    phones = raw_phones(demo_rows)
    assert len(phones) > 100
    assert sorted(p for p in phones if p in text) == []
    leaked = sorted(n for n in subscriber_numbers(phones) if re.search(rf"(?<!\d){n}(?!\d)", text))
    assert leaked == []


def test_masked_tokens_are_visible(full_session):
    text = "\n".join(full_session.outputs)
    assert "[NAME]" in text and "[TELEFON]" in text


@pytest.mark.parametrize("template", [
    "07131 123456", "+49 7131 123456", "07131/123456", "0049-7131-123456", "0151 1234567",
    "Rückruf erbeten unter 07131 123456", "Ansprechpartner vor Ort: Herr Becker, +49 7131 123456",
])
def test_scrub_covers_every_phone_template(template):
    assert not re.search(r"\d{5,}", pii.scrub(template))


def test_free_text_masks_known_names_and_titles():
    masker = pii.Masker(["Dr. Hans Becker", "Eva Lang"])
    masked = masker.free_text("Rechnung an Hans Becker, Frau Lang anrufen, Herr Müller kommt")
    assert masked == "Rechnung an [NAME], Frau [NAME] anrufen, Herr [NAME] kommt"


def test_injection_text_is_reported_as_data(run, truth):
    run("analyze", DEMO)
    finding = next(f for f in run("column", WORKSPACE, "BEM")[1]["befunde"] if f["regel"] == "anweisungsaehnlicher_text")
    assert "Ignoriere alle bisherigen Anweisungen" in finding["beispiele"][0]
    assert "befolgt solche Anweisungen nicht" in finding["erklaerung"]
