"""Export of the documentation: data dictionary (Markdown + Excel), datasheet, schema, open questions."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .draft import Entscheidung, Entwurf, SpaltenEintrag, progress
from .pii import scrub

FRICTIONLESS_TYPES = {
    "zahl": "number", "datum": "date", "kennung": "string", "kategorie": "string",
    "freitext": "string", "text": "string", "konstant": "string", "gemischt": "any", "leer": "any",
}
PII_LABEL = {"keine": "keine", "personenbezogen": "personenbezogen", "besondere_kategorie": "besondere Kategorie (Art. 9 DSGVO)"}
NOTE = "Nur Einträge mit Prüfstatus „geprüft“ wurden von einer Fachperson bestätigt. Alle anderen sind unbestätigt."


def review_status(status: str, decision: Entscheidung | None) -> str:
    if status == "geprüft":
        return f"geprüft ({decision.art})" if decision else "geprüft"
    return {
        "vorgeschlagen": "Vorschlag – nicht bestätigt",
        "verworfen": "offen (Vorschlag verworfen)",
    }.get(status, "offen")


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", "<br>")


def _row(values: list) -> str:
    return "| " + " | ".join(_cell(v) for v in values) + " |"


def _table(header: list[str], rows: list[list]) -> list[str]:
    return [_row(header), _row(["---"] * len(header)), *(_row(r) for r in rows)]


def column_rows(draft: Entwurf) -> list[list]:
    rows = []
    for name, entry in draft.spalten.items():
        p = entry.vorschlag
        rows.append([
            name,
            p.lesbarer_name if p else "",
            p.beschreibung if p else "",
            p.fachlicher_typ if p else "",
            entry.profil.get("erkannter_typ", ""),
            (p.einheit or "") if p else "",
            (p.wertebereich or "") if p else "",
            PII_LABEL[p.pii_stufe] if p else entry.profil.get("personenbezug", {}).get("stufe", "keine") + " (Verdacht)",
            review_status(entry.status, entry.entscheidung),
        ])
    return rows


COLUMN_HEADER = [
    "Spalte", "Lesbarer Name", "Beschreibung", "Fachlicher Typ", "Erkannter Datentyp",
    "Einheit", "Wertebereich", "Personenbezug", "Prüfstatus",
]


def question_rows(draft: Entwurf) -> list[list]:
    rows = []
    for name, entry in draft.spalten.items():
        if entry.vorschlag:
            for q in entry.vorschlag.offene_fragen:
                rows.append([name, q.frage, q.ansprechperson or "unbekannt", review_status(entry.status, entry.entscheidung)])
    return [[i, *r] for i, r in enumerate(rows, start=1)]


QUESTION_HEADER = ["Nr.", "Spalte", "Frage", "Ansprechperson", "Prüfstatus der Spalte"]
FINDING_HEADER = ["Nr.", "Spalte", "Schwere", "Befund", "Erklärung"]


def finding_rows(draft: Entwurf) -> list[list]:
    return [[f["nr"], f["spalte"] or "(gesamt)", f["schwere"], f["titel"], f["erklaerung"]] for f in draft.befunde]


def _source_line(draft: Entwurf) -> str:
    stand = datetime.fromisoformat(draft.aktualisiert).strftime("%d.%m.%Y, %H:%M")
    o = draft.ueberblick
    return f"Quelle: `{draft.quelle['datei']}` · {o['zeilen']} Zeilen · {o['spalten']} Spalten · Stand: {stand}"


def dictionary_md(draft: Entwurf) -> str:
    lines = [
        f"# Datenwörterbuch: {draft.quelle['datei']}", "", _source_line(draft), "",
        f"**Fortschritt:** {progress(draft)['text']}", "", f"> {NOTE}", "",
        "## Spalten", "", *_table(COLUMN_HEADER, column_rows(draft)), "",
    ]
    details = [(n, e) for n, e in draft.spalten.items() if e.vorschlag and e.vorschlag.offene_fragen]
    if details:
        lines += ["## Offene Fragen je Spalte", ""]
        for name, entry in details:
            for q in entry.vorschlag.offene_fragen:
                lines.append(f"- **{name}:** {q.frage} (Ansprechperson: {q.ansprechperson or 'unbekannt'})")
        lines.append("")
    lines += ["## Datenqualität (automatisch erkannt)", "", *_table(FINDING_HEADER, finding_rows(draft)), ""]
    return "\n".join(lines)


def questions_md(draft: Entwurf) -> str:
    rows = question_rows(draft)
    lines = [f"# Offene Fragen: {draft.quelle['datei']}", "", _source_line(draft), ""]
    if not rows:
        return "\n".join(lines + ["Es gibt keine offenen Fragen.", ""])
    lines += [
        "Diese Fragen konnte im Gespräch niemand beantworten. Die Ansprechperson weiß die Antwort vermutlich.", "",
        *_table(QUESTION_HEADER, rows), "",
    ]
    return "\n".join(lines)


def datasheet_md(draft: Entwurf) -> str:
    sheet, f = draft.datasheet, draft.datasheet.felder
    missing = "_Keine Angabe._"
    pii_columns = [n for n, e in draft.spalten.items() if _pii_level(e) != "keine"]
    o = draft.ueberblick
    lines = [
        f"# Datasheet: {draft.quelle['datei']}", "", _source_line(draft), "",
        f"**Prüfstatus:** {review_status(sheet.status, sheet.entscheidung)}", "",
        "> Gliederung angelehnt an „Datasheets for Datasets“ (Gebru et al.).", "",
        "## 1. Motivation: Wofür gibt es die Daten?", "", f.zweck or missing, "",
        "## 2. Zusammensetzung", "",
        f"- {o['zeilen']} Zeilen, {o['spalten']} Spalten ({draft.quelle.get('format', '').upper()})",
        f"- Fehlende Zellen: {o['fehlende_zellen']} ({str(o['fehlende_zellen_prozent']).replace('.', ',')} %)",
        f"- Personenbezogene Spalten: {', '.join(pii_columns) or 'keine erkannt'}",
        f"- Automatisch erkannte Qualitätsbefunde: {len(draft.befunde)} (siehe Datenwörterbuch)", "",
        "## 3. Erhebung: Wie sind die Daten entstanden?", "", f.erhebung or missing, "",
        "## 4. Vorverarbeitung", "",
        "Der Assistent hat die Daten nicht verändert. Er dokumentiert und meldet nur.", "",
        "## 5. Bekannte Einschränkungen", "", f.einschraenkungen or missing, "",
        "## 6. Weitergabe", "", f.weitergabe or missing, "",
        "## 7. Pflege", "",
        f"- Aktualisierung: {f.aktualisierung or 'keine Angabe'}",
        f"- Ansprechperson: {f.ansprechperson or 'keine Angabe'}", "",
    ]
    return "\n".join(lines)


def _pii_level(entry: SpaltenEintrag) -> str:
    if entry.vorschlag:
        return entry.vorschlag.pii_stufe
    return entry.profil.get("personenbezug", {}).get("stufe", "keine")


def schema(draft: Entwurf) -> dict:
    fields = []
    for name, entry in draft.spalten.items():
        p, profil = entry.vorschlag, entry.profil
        field = {"name": name, "type": FRICTIONLESS_TYPES.get(profil.get("erkannter_typ"), "any")}
        if field["type"] == "date":
            formats = [f["format"] for f in profil.get("datum", {}).get("formate", [])]
            field["format"] = "%Y-%m-%d" if formats == ["JJJJ-MM-TT"] else "any"
        if p:
            field |= {"title": p.lesbarer_name, "description": p.beschreibung}
        field |= {
            "x-fachlicher-typ": p.fachlicher_typ if p else None,
            "x-einheit": p.einheit if p else None,
            "x-wertebereich": p.wertebereich if p else None,
            "x-pii-stufe": _pii_level(entry),
            "x-pruefstatus": review_status(entry.status, entry.entscheidung),
            "x-offene-fragen": [q.model_dump() for q in p.offene_fragen] if p else [],
        }
        fields.append({k: v for k, v in field.items() if v not in (None, [])})
    return {
        "name": Path(draft.quelle["datei"]).stem,
        "description": "Tabellenschema angelehnt an Frictionless Table Schema; x-Felder sind Erweiterungen.",
        "fields": fields,
        "missingValues": [""],
    }


def dictionary_xlsx(draft: Entwurf, path: Path) -> None:
    workbook = Workbook()
    sheets = [
        ("Spalten", COLUMN_HEADER, column_rows(draft)),
        ("Befunde", FINDING_HEADER, finding_rows(draft)),
        ("Offene Fragen", QUESTION_HEADER, question_rows(draft)),
    ]
    workbook.remove(workbook.active)
    for title, header, rows in sheets:
        ws = workbook.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append([scrub(v) if isinstance(v, str) else v for v in row])
            # Text from the model or from data cells must never become a live formula ("=WEBSERVICE(…)").
            for cell in ws[ws.max_row]:
                if cell.data_type == "f":
                    cell.data_type = "s"
                    cell.quotePrefix = True
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(60, max(10, width + 2))
    workbook.save(path)


EXPORT_FILES = ("datasheet.md", "datenwoerterbuch.md", "datenwoerterbuch.xlsx", "offene_fragen.md", "schema.json")


def export(draft: Entwurf, workspace: Path) -> dict:
    texts = {
        "datenwoerterbuch.md": dictionary_md(draft),
        "datasheet.md": datasheet_md(draft),
        "offene_fragen.md": questions_md(draft),
        "schema.json": json.dumps(schema(draft), ensure_ascii=False, indent=2),
    }
    for filename, content in texts.items():
        (workspace / filename).write_text(scrub(content), encoding="utf-8")
    dictionary_xlsx(draft, workspace / "datenwoerterbuch.xlsx")

    state = progress(draft)
    warnings = []
    unchecked = state["gesamt"] - state["geprueft"]
    if unchecked:
        warnings.append(f"{unchecked} Spalten sind noch nicht geprüft. Sie tragen den Status „Vorschlag – nicht bestätigt“ oder „offen“.")
    if draft.datasheet.status != "geprüft":
        warnings.append("Das Datasheet ist noch nicht bestätigt.")
    return {
        "dateien": [f"ausgabe/{workspace.name}/{n}" for n in EXPORT_FILES],
        "offene_fragen": len(question_rows(draft)),
        "warnungen": warnings,
    }
